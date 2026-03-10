"""Grade tracker — logs grades to an Excel spreadsheet with history."""

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .config import GRADES_FILE
from .scraper import GradeEntry


# ── Colour palette ───────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="EEF2F7")
DATE_FILL = PatternFill("solid", fgColor="D9E8FB")


# ── Workbook helpers ─────────────────────────────────────────────────────────

def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


# ── Main functions ───────────────────────────────────────────────────────────

def update_grade_spreadsheet(
    grades: list[GradeEntry],
    output_path: Path = GRADES_FILE,
) -> Path:
    """
    Write grades to an Excel workbook.

    Structure:
      - Sheet 'Overview': one row per grade item with latest value
      - Sheet 'History': timestamped log — appended on every run
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now()
    run_stamp = now.strftime("%Y-%m-%d %H:%M")

    if output_path.exists():
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    # ── Overview sheet ───────────────────────────────────────────────────────
    if "Overview" in wb.sheetnames:
        del wb["Overview"]
    ws_ov = wb.create_sheet("Overview", 0)
    _header_row(ws_ov, ["Course", "Item", "Grade", "Percentage", "Last Updated"])

    for i, g in enumerate(grades, 2):
        ws_ov.append([g.course, g.item, g.grade, g.percentage, run_stamp])
        if i % 2 == 0:
            for col in range(1, 6):
                ws_ov.cell(row=i, column=col).fill = ALT_FILL

    _auto_width(ws_ov)

    # ── History sheet ────────────────────────────────────────────────────────
    if "History" not in wb.sheetnames:
        ws_hist = wb.create_sheet("History")
        _header_row(ws_hist, ["Timestamp", "Course", "Item", "Grade", "Percentage"])
    else:
        ws_hist = wb["History"]

    for g in grades:
        ws_hist.append([run_stamp, g.course, g.item, g.grade, g.percentage])

    _auto_width(ws_hist)

    wb.save(output_path)
    print(f"[grades] Saved {len(grades)} grade entries to {output_path}")
    return output_path


def print_grade_summary(grades: list[GradeEntry]) -> None:
    """Print a quick grade summary grouped by course."""
    from rich.console import Console
    from rich.table import Table
    from rich import box

    console = Console()
    courses = sorted({g.course for g in grades})

    for course in courses:
        course_grades = [g for g in grades if g.course == course]
        table = Table(title=course, box=box.SIMPLE_HEAD)
        table.add_column("Item", style="white")
        table.add_column("Grade", style="cyan", justify="right")
        table.add_column("%", style="green", justify="right")

        for g in course_grades:
            table.add_row(g.item, g.grade, g.percentage)

        console.print(table)
        console.print()
